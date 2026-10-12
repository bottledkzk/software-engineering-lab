#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""md2docx.py —— 把本项目 Markdown 文档转换为 DOCX（python-docx）

用法:
    python md2docx.py <input.md> <output.docx>

支持的 Markdown 子集：标题(#~####)、段落、项目符号/有序列表、管道表格、
引用块(>)、围栏代码块(```)、水平线(---)、行内 **粗体** / `代码` / [链接](url)。
"""
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "微软雅黑"
MONO_FONT = "Consolas"
BODY_SIZE = Pt(10.5)
TABLE_SIZE = Pt(9)


# ---------- 底层工具 ----------
def set_font(run, name=BODY_FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:eastAsia"), BODY_FONT)
    rfonts.set(qn("w:ascii"), BODY_FONT)
    rfonts.set(qn("w:hAnsi"), BODY_FONT)
    rpr.append(rfonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0645AD")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)\s]+\))")


def add_inline(paragraph, text, size=None, base_bold=False):
    """处理行内格式：**粗体**、`代码`、[文本](链接)"""
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            run = paragraph.add_run(chunk[2:-2])
            set_font(run, BODY_FONT, size, bold=True)
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            run = paragraph.add_run(chunk[1:-1])
            set_font(run, MONO_FONT, size or Pt(9.5))
        else:
            m = re.fullmatch(r"\[([^\]]+)\]\(([^)\s]+)\)", chunk)
            if m:
                add_hyperlink(paragraph, m.group(2), m.group(1))
            else:
                run = paragraph.add_run(chunk)
                set_font(run, BODY_FONT, size, bold=True if base_bold else None)


def shade(cell, hex_color="EAEAEA"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_page_number_footer(section):
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_font(run, BODY_FONT, Pt(9))
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def horizontal_rule(document):
    paragraph = document.add_paragraph()
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "B0B0B0")
    borders.append(bottom)
    p_pr.append(borders)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)


# ---------- 文档初始化 ----------
def init_document():
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    for level, size in ((1, Pt(15)), (2, Pt(13)), (3, Pt(11.5)), (4, Pt(11))):
        style = document.styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style.font.size = size
        style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    add_page_number_footer(section)
    return document


# ---------- 解析与写入 ----------
def split_row(line):
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    return [c.strip() for c in text.split("|")]


def is_table_sep(line):
    return bool(re.fullmatch(r"\s*\|[\s:|-]+\|\s*", line or "")) and "-" in line


def convert(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as handle:
        lines = handle.read().replace("\r\n", "\n").split("\n")

    document = init_document()
    i = 0
    in_code = False
    code_buffer = []
    para_buffer = []
    list_mode = None

    def flush_paragraph():
        nonlocal para_buffer
        if para_buffer:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            add_inline(paragraph, " ".join(para_buffer))
            para_buffer = []

    def close_list():
        nonlocal list_mode
        list_mode = None

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            if not in_code:
                in_code, code_buffer = True, []
            else:
                in_code = False
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.4)
                paragraph.paragraph_format.space_after = Pt(6)
                run = paragraph.add_run("\n".join(code_buffer))
                set_font(run, MONO_FONT, Pt(9))
                shade_paragraph(paragraph)
            i += 1
            continue
        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        # 表格
        if stripped.startswith("|") and is_table_sep(lines[i + 1] if i + 1 < len(lines) else ""):
            flush_paragraph()
            close_list()
            header = split_row(line)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            table = document.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for idx, text in enumerate(header):
                cell = table.rows[0].cells[idx]
                cell.text = ""
                add_inline(cell.paragraphs[0], text, size=TABLE_SIZE, base_bold=True)
                shade(cell)
            for row in rows:
                cells = table.add_row().cells
                for idx in range(len(header)):
                    value = row[idx] if idx < len(row) else ""
                    cells[idx].text = ""
                    add_inline(cells[idx].paragraphs[0], value, size=TABLE_SIZE)
            document.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # 标题
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            flush_paragraph()
            close_list()
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, text, size=Pt(18), base_bold=True)
            else:
                paragraph = document.add_heading(level=min(level - 1, 4))
                paragraph.clear()
                add_inline(paragraph, text, base_bold=True)
            i += 1
            continue

        # 水平线
        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            flush_paragraph()
            close_list()
            horizontal_rule(document)
            i += 1
            continue

        # 引用块
        if stripped.startswith(">"):
            flush_paragraph()
            close_list()
            buffer = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buffer.append(lines[i].strip().lstrip(">").strip())
                i += 1
            for item in buffer:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.5)
                paragraph.paragraph_format.space_after = Pt(2)
                if item:
                    add_inline(paragraph, item, size=Pt(9.5))
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
            continue

        # 列表
        bullet = re.match(r"^\s*[-*]\s+(.*)$", line)
        numbered = re.match(r"^\s*\d+[.)]\s+(.*)$", line)
        if bullet or numbered:
            flush_paragraph()
            style = "List Bullet" if bullet else "List Number"
            if list_mode != style:
                list_mode = style
            paragraph = document.add_paragraph(style=style)
            paragraph.paragraph_format.space_after = Pt(2)
            add_inline(paragraph, (bullet or numbered).group(1))
            i += 1
            continue

        # 空行
        if stripped == "":
            flush_paragraph()
            close_list()
            i += 1
            continue

        para_buffer.append(stripped)
        i += 1

    flush_paragraph()
    document.save(docx_path)
    return docx_path


def shade_paragraph(paragraph, hex_color="F5F5F5"):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    p_pr.append(shd)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("usage: python md2docx.py <input.md> <output.docx>")
        sys.exit(1)
    out = convert(sys.argv[1], sys.argv[2])
    print(f"OK {out}")
